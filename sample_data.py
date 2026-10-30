"""
Generador de archivos de muestra SENA en formato Excel (.xlsx) usando openpyxl (100% pure Python).
Programa: TECNÓLOGO EN ANÁLISIS Y DESARROLLO DE SOFTWARE (ADSO) - Ficha 2714589
"""

import os
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

NOMBRES = [
    "ALEJANDRO", "CAROLINA", "DANIEL", "VALENTINA", "MATEO",
    "CAMILA", "SEBASTIAN", "MARIANA", "JUAN DAVID", "SARA",
    "NICOLAS", "ANDREA", "SANTIAGO", "PAULA", "FELIPE",
    "DANIELA", "JULIAN", "MANUELA", "GABRIEL", "ISABELLA",
    "CRISTIAN", "NATALIA", "KEVIN", "LAURA", "BRAYAN"
]

APELLIDOS = [
    "RODRIGUEZ PEREZ", "GOMEZ MARTINEZ", "GONZALEZ LOPEZ", "GARCIA SANCHEZ",
    "HERNANDEZ DIAZ", "MORALES CASTRO", "RAMIREZ TORRES", "VARGAS ROJAS",
    "JIMENEZ SILVA", "CASTILLO ORTIZ", "GUTIERREZ RUIZ", "MORENO VEGA",
    "MENDOZA RIOS", "ALVAREZ CARDONA", "ROMERO DUQUE", "SUAREZ MEJIA",
    "HERRERA GUZMAN", "MEDINA OROZCO", "AGUILAR SALAZAR", "CORREA RESTREPO",
    "PARRA CARDENAS", "VALENCIA BUSTAMANTE", "QUINTERO HOYOS", "LONDOÑO MONTOYA", "MONTOYA RESTREPO"
]

PROGRAMA = "TECNÓLOGO EN ANÁLISIS Y DESARROLLO DE SOFTWARE"
FICHA = "2714589"

COMPETENCIAS_RAPS = [
    {
        "competencia": "220501096 - Estructurar la base de datos de acuerdo con los requerimientos del cliente",
        "raps": [
            "RAP 01: Diseñar el modelo entidad-relación conceptual y lógico según requisitos.",
            "RAP 02: Implementar la base de datos relacional en motores SQL según diseño.",
            "RAP 03: Construir consultas, vistas y procedimientos almacenados para la persistencia."
        ]
    },
    {
        "competencia": "220501095 - Construir el software según especificaciones de diseño y estándares",
        "raps": [
            "RAP 04: Desarrollar interfaces web responsivas utilizando HTML, CSS y JavaScript.",
            "RAP 05: Implementar lógica de negocio y controladores del lado servidor.",
            "RAP 06: Integrar servicios web y APIs REST cumpliendo protocolos de seguridad."
        ]
    },
    {
        "competencia": "240201528 - Razonar cuantitativamente frente a situaciones del contexto laboral",
        "raps": [
            "RAP 07: Plantear problemas aritméticos y algebraicos a partir de situaciones laborales.",
            "RAP 08: Resolver problemas de lógica matemática y algorítmica aplicada al desarrollo."
        ]
    },
    {
        "competencia": "240202501 - Interactuar en lengua inglesa de forma oral y escrita",
        "raps": [
            "RAP 09: Comprender textos técnicos y manuales de desarrollo en idioma inglés.",
            "RAP 10: Redactar especificaciones y comentarios de código en lengua inglesa."
        ]
    }
]


def generate_sample_excel(output_path: str = "data/muestra_sena_juicios_evaluativos.xlsx"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    random.seed(42)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Juicios Evaluativos"

    # Encabezados oficiales de SofiaPlus / Reporte de Seguimiento
    headers = [
        "Número Ficha",
        "Programa de Formación",
        "Número Documento",
        "Nombre Aprendiz",
        "Estado del Aprendiz",
        "Competencia",
        "Resultado de Aprendizaje",
        "Juicio Evaluativo",
        "Instructor Evaluador"
    ]
    ws.append(headers)

    # Estilos encabezado SENA Verde (#39A900)
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="39A900", end_color="39A900", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Crear aprendices
    num_aprendices = 25
    aprendices = []
    for i in range(num_aprendices):
        doc = f"{1000000000 + (i * 38271) + 12345}"
        nom = f"{NOMBRES[i]} {APELLIDOS[i]}"
        if i < 15:
            perfil = "alto"
            estado = "EN FORMACIÓN"
        elif i < 21:
            perfil = "medio"
            estado = "EN FORMACIÓN"
        elif i < 23:
            perfil = "rezago"
            estado = "EN FORMACIÓN"
        elif i == 23:
            perfil = "desercion"
            estado = "CANCELADO"
        else:
            perfil = "desercion"
            estado = "RETIRO VOLUNTARIO"
        aprendices.append((doc, nom, estado, perfil))

    for comp_idx, comp_data in enumerate(COMPETENCIAS_RAPS):
        comp_name = comp_data["competencia"]
        instructor = "ING. CARLOS ARTURO LONDOÑO" if comp_idx < 2 else ("LIC. ADRIANA VALENCIA" if comp_idx == 2 else "LIC. JOHN FREDY OSPINA")

        for rap_idx, rap_name in enumerate(comp_data["raps"]):
            is_early_rap = comp_idx == 0 or (comp_idx == 1 and rap_idx == 0) or (comp_idx == 2 and rap_idx == 0)
            is_mid_rap = (comp_idx == 1 and rap_idx == 1) or (comp_idx == 3 and rap_idx == 0)

            for doc, nom, estado, perfil in aprendices:
                if estado in ["CANCELADO", "RETIRO VOLUNTARIO"]:
                    juicio = "A" if (is_early_rap and rap_idx == 0 and random.random() > 0.5) else "POR EVALUAR"
                elif perfil == "alto":
                    if is_early_rap:
                        juicio = "A"
                    elif is_mid_rap:
                        juicio = "A" if random.random() < 0.88 else "POR EVALUAR"
                    else:
                        juicio = "A" if random.random() < 0.45 else "POR EVALUAR"
                elif perfil == "medio":
                    if is_early_rap:
                        juicio = "A" if random.random() < 0.8 else "D"
                    elif is_mid_rap:
                        r = random.random()
                        juicio = "A" if r < 0.55 else ("D" if r < 0.75 else "POR EVALUAR")
                    else:
                        juicio = "A" if random.random() < 0.25 else "POR EVALUAR"
                else:  # rezago
                    if is_early_rap:
                        juicio = "A" if random.random() < 0.3 else "D"
                    elif is_mid_rap:
                        juicio = "D" if random.random() < 0.6 else "POR EVALUAR"
                    else:
                        juicio = "POR EVALUAR"

                ws.append([
                    FICHA,
                    PROGRAMA,
                    doc,
                    nom,
                    estado,
                    comp_name,
                    rap_name,
                    juicio,
                    instructor
                ])

    # Autoajuste de ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    p = generate_sample_excel()
    print(f"Muestra SENA generada con éxito en: {p}")
