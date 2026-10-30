"""
Generador de Informes Ejecutivos en Excel para el SENA.
Crea un libro de trabajo estilizado con pestaña de resumen, semáforo de RAPs,
aprendices en riesgo y matriz de seguimiento.
"""

import io
from typing import Any, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_excel_report(analysis: Dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # PESTAÑA 1: RESUMEN EJECUTIVO
    # -------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"
    ws_resumen.views.sheetView[0].showGridLines = True

    # Paleta SENA
    sena_green = "39A900"
    sena_dark = "00324D"
    card_bg = "F4F9F4"
    alert_red = "FF4D4D"
    alert_yellow = "FFB020"

    title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    section_font = Font(name="Segoe UI", size=12, bold=True, color=sena_dark)
    kpi_num_font = Font(name="Segoe UI", size=20, bold=True, color=sena_green)
    kpi_label_font = Font(name="Segoe UI", size=9, bold=True, color="555555")

    # Banner Superior
    for col in range(1, 7):
        for row in range(1, 4):
            c = ws_resumen.cell(row=row, column=col)
            c.fill = PatternFill(start_color=sena_dark, end_color=sena_dark, fill_type="solid")

    ws_resumen.merge_cells("A1:F2")
    banner_cell = ws_resumen.cell(row=1, column=1)
    banner_cell.value = "SENA - INFORME EJECUTIVO DE AVANCE REAL FRENTE A RAPS"
    banner_cell.font = title_font
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_resumen.merge_cells("A3:F3")
    meta_cell = ws_resumen.cell(row=3, column=1)
    meta = analysis["meta"]
    meta_cell.value = f"Ficha: {meta['ficha']}  |  Programa: {meta['programa']}  |  Aprendices Analizados: {meta['total_aprendices_analizados']}"
    meta_cell.font = subtitle_font
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")

    # KPIs Principales
    kpis = analysis["kpis"]
    kpi_data = [
        ("NIVEL DE AVANCE REAL", f"{kpis['porcentaje_avance_real']}%", 5),
        ("APROBACIÓN SOBRE EVALUADOS", f"{kpis['porcentaje_aprobacion_evaluados']}%", 5),
        ("COBERTURA EVALUATIVA", f"{kpis['porcentaje_cobertura_evaluativa']}%", 5),
        ("RAPS EN REZAGO CRÍTICO", str(kpis['raps_criticos']), 8),
        ("APRENDICES EN ALTO RIESGO", str(kpis['aprendices_riesgo']), 8),
        ("APRENDICES AL DÍA", str(kpis['aprendices_al_dia']), 8),
    ]

    # Render tarjetas KPI
    row_offset = 5
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    ws_resumen.cell(row=row_offset, column=1, value="INDICADORES CLAVE DEL GRUPO").font = section_font
    row_offset += 2

    # Fila 1 KPIs
    for idx, (label, val, _) in enumerate(kpi_data[:3]):
        c_val = ws_resumen.cell(row=row_offset, column=idx * 2 + 1, value=val)
        c_val.font = kpi_num_font
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        
        c_lbl = ws_resumen.cell(row=row_offset + 1, column=idx * 2 + 1, value=label)
        c_lbl.font = kpi_label_font
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")

    row_offset += 3
    # Fila 2 KPIs
    for idx, (label, val, _) in enumerate(kpi_data[3:]):
        c_val = ws_resumen.cell(row=row_offset, column=idx * 2 + 1, value=val)
        c_val.font = Font(name="Segoe UI", size=20, bold=True, color=alert_red if "CRÍTICO" in label or "RIESGO" in label else sena_green)
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        
        c_lbl = ws_resumen.cell(row=row_offset + 1, column=idx * 2 + 1, value=label)
        c_lbl.font = kpi_label_font
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")

    # Diagnóstico y Recomendaciones
    row_offset += 4
    ws_resumen.cell(row=row_offset, column=1, value="DIAGNÓSTICO Y PLAN DE ACCIÓN SUGERIDO").font = section_font
    row_offset += 1

    diag_text = [
        f"1. Cobertura del Grupo: El grupo presenta un avance real del {kpis['porcentaje_avance_real']}%. Se han registrado {kpis['total_aprobados']} aprobaciones de {kpis['total_evaluaciones_posibles']} evaluaciones esperadas.",
        f"2. Resultados Críticos: Se identificaron {kpis['raps_criticos']} RAPs con avance inferior al {meta['umbral_amarillo']}%. Se recomienda concertar planes de refuerzo inmediatos.",
        f"3. Situación de Aprendices: {kpis['aprendices_riesgo']} aprendices presentan rezago severo. Se sugiere remitir a Comité de Evaluación y Seguimiento según Reglamento del Aprendiz.",
        f"4. Aprendices al Día: {kpis['aprendices_al_dia']} aprendices cumplen con la totalidad de los juicios programados a la fecha."
    ]
    for d in diag_text:
        ws_resumen.merge_cells(start_row=row_offset, start_column=1, end_row=row_offset, end_column=6)
        ws_resumen.cell(row=row_offset, column=1, value=d).font = Font(name="Segoe UI", size=10)
        row_offset += 1

    # -------------------------------------------------------------
    # PESTAÑA 2: SEMÁFORO DE RAPS
    # -------------------------------------------------------------
    ws_raps = wb.create_sheet(title="Semáforo de RAPs")
    ws_raps.views.sheetView[0].showGridLines = True
    
    headers_raps = [
        "Competencia", "Resultado de Aprendizaje (RAP)", "Instructor",
        "Aprobados", "Deficientes", "Por Evaluar", "% Avance Real", "Estado"
    ]
    ws_raps.append(headers_raps)

    header_fill = PatternFill(start_color=sena_dark, end_color=sena_dark, fill_type="solid")
    header_font_white = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers_raps) + 1):
        c = ws_raps.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font_white
        c.alignment = Alignment(horizontal="center", vertical="center")

    fill_green = PatternFill(start_color="E8F8EA", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFF8E7", fill_type="solid")
    fill_red = PatternFill(start_color="FDECEB", fill_type="solid")

    for rap in analysis["raps"]:
        row_data = [
            rap["competencia"],
            rap["rap"],
            rap["instructor"],
            rap["aprobados"],
            rap["deficientes"],
            rap["pendientes"],
            f"{rap['porcentaje_aprobado']}%",
            rap["semaforo"]
        ]
        ws_raps.append(row_data)
        curr_row = ws_raps.max_row
        
        # Color según estado
        row_fill = fill_green if rap["categoria_cod"] == "optimo" else (fill_yellow if rap["categoria_cod"] == "en_avance" else fill_red)
        for col_idx in range(1, len(headers_raps) + 1):
            cell = ws_raps.cell(row=curr_row, column=col_idx)
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [4, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # -------------------------------------------------------------
    # PESTAÑA 3: APRENDICES Y RIESGO ACADÉMICO
    # -------------------------------------------------------------
    ws_apr = wb.create_sheet(title="Seguimiento de Aprendices")
    ws_apr.views.sheetView[0].showGridLines = True

    headers_apr = [
        "Documento", "Nombre Aprendiz", "Estado Sofía", "Aprobados",
        "Deficientes", "Pendientes", "% Avance", "Situación Académica", "RAPs Deficientes (A Recuperar)"
    ]
    ws_apr.append(headers_apr)
    for col_idx in range(1, len(headers_apr) + 1):
        c = ws_apr.cell(row=1, column=col_idx)
        c.fill = PatternFill(start_color=sena_green, end_color=sena_green, fill_type="solid")
        c.font = header_font_white
        c.alignment = Alignment(horizontal="center", vertical="center")

    for apr in analysis["aprendices"]:
        def_str = "; ".join(apr["raps_deficientes_lista"]) if apr["raps_deficientes_lista"] else "Sin juicios deficientes"
        ws_apr.append([
            apr["documento"],
            apr["nombre"],
            apr["estado_sofia"],
            apr["raps_aprobados"],
            apr["raps_deficientes"],
            apr["raps_pendientes"],
            f"{apr['porcentaje_avance']}%",
            apr["situacion"],
            def_str
        ])
        curr_row = ws_apr.max_row
        row_fill = fill_green if apr["situacion_cod"] == "al_dia" else (fill_yellow if apr["situacion_cod"] == "observacion" else fill_red)
        for col_idx in range(1, len(headers_apr) + 1):
            cell = ws_apr.cell(row=curr_row, column=col_idx)
            cell.fill = row_fill
            cell.border = thin_border
            if col_idx in [1, 3, 4, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # -------------------------------------------------------------
    # PESTAÑA 4: MATRIZ APRENDIZ X RAP
    # -------------------------------------------------------------
    ws_mat = wb.create_sheet(title="Matriz Aprendiz vs RAP")
    ws_mat.views.sheetView[0].showGridLines = True

    matriz_data = analysis["matriz"]
    headers_mat = ["Documento", "Nombre Aprendiz"] + [f"RAP {i+1}" for i in range(len(matriz_data["raps"]))]
    ws_mat.append(headers_mat)
    for col_idx in range(1, len(headers_mat) + 1):
        c = ws_mat.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font_white
        c.alignment = Alignment(horizontal="center", vertical="center")

    fill_a = PatternFill(start_color="39A900", end_color="39A900", fill_type="solid")
    font_a = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_d = PatternFill(start_color="E53E3E", end_color="E53E3E", fill_type="solid")
    font_d = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_p = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    font_p = Font(name="Segoe UI", size=10, color="718096")

    for f in matriz_data["filas"]:
        row = [f["documento"], f["nombre"]]
        for r_name in matriz_data["raps"]:
            row.append(f["juicios"].get(r_name, "P"))
        ws_mat.append(row)
        curr_row = ws_mat.max_row
        for col_idx in range(3, len(row) + 1):
            cell = ws_mat.cell(row=curr_row, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            val = str(cell.value)
            if val == "A":
                cell.fill = fill_a
                cell.font = font_a
            elif val == "D":
                cell.fill = fill_d
                cell.font = font_d
            else:
                cell.fill = fill_p
                cell.font = font_p

    # Ajuste automático de anchos para todas las hojas
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
